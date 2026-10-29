"""Revisa la planilla de productos y las fotos antes de subirlas a Shopify.

Uso (desde la carpeta del proyecto):
    python catalogo/revisar_catalogo.py

Avisa, producto por producto, lo que falta según el estándar de carga
(ver catalogo/README.md). No cambia nada: solo lee.
"""
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

CARPETA = Path(__file__).resolve().parent
PLANILLA = CARPETA / "productos.xlsx"
FOTOS = CARPETA / "fotos"
MINIMO_FOTOS = 3
EXTENSIONES = {".jpg", ".jpeg", ".png", ".webp"}
OBLIGATORIAS = ["nombre", "coleccion", "categoria", "material", "tamano", "descripcion", "estado"]
CON_LISTA = ["coleccion", "categoria", "material", "estado"]
FORMATO_CODIGO = re.compile(r"^[a-z0-9]+(-[a-z0-9]+)*$")


def leer_listas(wb):
    hoja = wb["Listas"]
    listas = {}
    for col in hoja.iter_cols(min_row=1):
        nombre = col[0].value
        if nombre:
            listas[nombre] = {str(c.value).strip() for c in col[1:] if c.value not in (None, "")}
    return listas


def leer_productos(wb):
    hoja = wb["Productos"]
    encabezados = [str(c.value or "").replace("*", "").strip() for c in hoja[1]]
    productos = {}
    for n, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        datos = {k: (v.strip() if isinstance(v, str) else v) for k, v in zip(encabezados, fila)}
        if not any(v not in (None, "") for v in datos.values()):
            continue
        codigo = datos.get("codigo")
        productos.setdefault(codigo, []).append((n, datos))
    return productos


def fotos_de(codigo, fotos=FOTOS):
    carpeta = fotos / str(codigo)
    return sorted(p for p in carpeta.glob("*") if p.suffix.lower() in EXTENSIONES) if carpeta.is_dir() else []


def revisar_producto(codigo, filas, listas, fotos=FOTOS):
    problemas = []
    n, principal = filas[0]
    if not codigo:
        return [f"fila {n}: falta el código"]
    if not FORMATO_CODIGO.match(str(codigo)):
        problemas.append("el código debe ir en minúsculas, sin tildes ni espacios, separado con guiones")
    for campo in OBLIGATORIAS:
        if principal.get(campo) in (None, ""):
            problemas.append(f"falta {campo} (fila {n})")
    for campo in CON_LISTA:
        valor = principal.get(campo)
        if valor not in (None, "") and str(valor) not in listas.get(campo, set()):
            problemas.append(f"{campo} «{valor}» no está en la hoja Listas")
    con_opcion = [f for f in filas if f[1].get("opcion") not in (None, "")]
    if con_opcion and len(con_opcion) != len(filas):
        problemas.append("si una versión tiene opción, todas sus filas deben tenerla")
    for fila_n, datos in filas:
        precio = datos.get("precio_clp")
        if precio in (None, ""):
            problemas.append(f"falta precio_clp (fila {fila_n})")
        elif not isinstance(precio, (int, float)) or precio <= 0:
            problemas.append(f"precio_clp no válido en la fila {fila_n}: {precio}")
        if fila_n != n and datos.get("nombre") not in (None, ""):
            problemas.append(f"el código se repite en la fila {fila_n} con otro nombre: si es otro producto, usa otro código")
    cantidad = len(fotos_de(codigo, fotos))
    if cantidad < MINIMO_FOTOS:
        problemas.append(f"tiene {cantidad} foto(s) en catalogo/fotos/{codigo}/; se piden al menos {MINIMO_FOTOS}")
    return problemas


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if not PLANILLA.exists():
        print(f"No encuentro la planilla: {PLANILLA}")
        return 2
    wb = load_workbook(PLANILLA, data_only=True)
    listas = leer_listas(wb)
    productos = leer_productos(wb)
    con_problemas = 0
    for codigo, filas in productos.items():
        problemas = revisar_producto(codigo, filas, listas)
        nombre = filas[0][1].get("nombre") or "(sin nombre)"
        if problemas:
            con_problemas += 1
            print(f"✗ {codigo}: {nombre}")
            for p in problemas:
                print(f"    - {p}")
        else:
            print(f"✓ {codigo}: {nombre}")
    sobrantes = [p.name for p in FOTOS.iterdir() if p.is_dir() and p.name not in productos] if FOTOS.is_dir() else []
    for s in sobrantes:
        print(f"! La carpeta catalogo/fotos/{s}/ no corresponde a ningún código de la planilla")
    print(f"\n{len(productos)} productos revisados: {len(productos) - con_problemas} listos, {con_problemas} con pendientes.")
    return 1 if con_problemas or sobrantes else 0


if __name__ == "__main__":
    sys.exit(main())
